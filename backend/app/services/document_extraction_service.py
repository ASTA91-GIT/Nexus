import os
import json
import pandas as pd
import fitz  # PyMuPDF

import asyncio

async def extract_text_from_file(file_path: str, filename: str) -> str:
    """
    Extracts text from the given file based on its extension.
    Supported: PDF, TXT, CSV, JSON, XLSX, DOCX (basic).
    """
    def _extract_sync():
        ext = os.path.splitext(filename)[1].lower()
        
        try:
            if ext == ".pdf":
                return extract_text_from_pdf(file_path)
            elif ext == ".txt" or ext == ".md" or ext == ".csv":
                return extract_text_from_txt(file_path)
            elif ext == ".json":
                return extract_text_from_json(file_path)
            elif ext in [".xls", ".xlsx"]:
                return extract_text_from_excel(file_path)
            elif ext == ".docx":
                return extract_text_from_docx(file_path)
            elif ext in [".webm", ".mp4", ".mp3", ".png", ".jpg", ".jpeg"]:
                return "UNSUPPORTED_MEDIA"
            else:
                # Fallback for unsupported or unknown text files
                try:
                    return extract_text_from_txt(file_path)
                except Exception:
                    return f"Unsupported file type for text extraction: {ext}"
        except Exception as e:
            return f"Error extracting content: {str(e)}"
            
    return await asyncio.to_thread(_extract_sync)

def extract_text_from_pdf(file_path: str) -> str:
    doc = fitz.open(file_path)
    text = ""
    for page in doc:
        text += page.get_text("text") + "\n"
    return text.strip()

def extract_text_from_txt(file_path: str) -> str:
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read().strip()

def extract_text_from_json(file_path: str) -> str:
    with open(file_path, "r", encoding="utf-8") as f:
        data = json.load(f)
        return json.dumps(data, indent=2)

def extract_text_from_excel(file_path: str) -> str:
    df = pd.read_excel(file_path)
    return df.to_string()

def extract_text_from_docx(file_path: str) -> str:
    try:
        import docx
    except ImportError:
        return "python-docx not installed"
        
    doc = docx.Document(file_path)
    text = []
    for para in doc.paragraphs:
        if para.text.strip():
            text.append(para.text.strip())
            
    for table in doc.tables:
        for row in table.rows:
            row_text = []
            for cell in row.cells:
                if cell.text.strip():
                    row_text.append(cell.text.strip())
            if row_text:
                text.append(" | ".join(row_text))
                
    return "\n".join(text).strip()
