import asyncio
import os
from docx import Document
from app.core.database import get_database
from app.api.routes.ingestion import process_file_background

async def create_test_docx(filename, doc_type):
    doc = Document()
    if doc_type == 'A':
        doc.add_heading('Case: 2008 Mumbai Attacks', 0)
        doc.add_paragraph('This is a test document.')
        table = doc.add_table(rows=2, cols=3)
        table.cell(0, 0).text = 'ID'
        table.cell(0, 1).text = 'Type'
        table.cell(0, 2).text = 'Name'
        table.cell(1, 0).text = 'E1'
        table.cell(1, 1).text = 'PERSON'
        table.cell(1, 2).text = 'Ajmal Amir Kasab'
        doc.add_paragraph('Kasab was associated with Lashkar-e-Taiba.')
    else:
        doc.add_heading('Case: Operation Cyber Storm', 0)
        doc.add_paragraph('Investigation into a corporate hack.')
        table = doc.add_table(rows=2, cols=3)
        table.cell(0, 0).text = 'ID'
        table.cell(0, 1).text = 'Type'
        table.cell(0, 2).text = 'Name'
        table.cell(1, 0).text = 'E1'
        table.cell(1, 1).text = 'ORGANIZATION'
        table.cell(1, 2).text = 'Cyberdyne Systems'
        doc.add_paragraph('Cyberdyne Systems developed Skynet in 1997.')
    
    doc.save(filename)

async def main():
    await create_test_docx("testA.docx", "A")
    await create_test_docx("testB.docx", "B")
    
    from dotenv import load_dotenv
    load_dotenv()
    
    from motor.motor_asyncio import AsyncIOMotorClient
    mongo_uri = os.getenv('MONGODB_URI', 'mongodb://localhost:27017')
    client = AsyncIOMotorClient(mongo_uri)
    db = client['nexus_db']
    
    current_user = {"email": "test@nexus.gov"}
    
    print("--- RUNNING TEST A ---")
    case_a_id = f'test_case_A_{uuid.uuid4().hex[:8]}'
    await db["cases"].insert_one({"_id": case_a_id, "name": "Test Case A"})
    ev_a_result = await db["evidence"].insert_one({"case_id": case_a_id, "title": "testA.docx", "source_type": "DOCX"})
    ev_a_id = str(ev_a_result.inserted_id)
    
    with open("testA.docx", "rb") as f:
        content_A = f.read()
        
    await process_file_background(db, case_a_id, content_A, "testA.docx", current_user, ev_a_id)
    
    ents_a = await db["entities"].find({"case_id": case_a_id}).to_list(length=100)
    rels_a = await db["relationships"].find({"case_id": case_a_id}).to_list(length=100)
    ev_a = await db["evidence"].find_one({"_id": ev_a_result.inserted_id})
    print(f"Text A Length: {len(ev_a.get('raw_content', ''))}")
    print(f"Entities in A: {len(ents_a)}")
    for e in ents_a:
        print(f" - {e['name']} ({e['type']})")
    print(f"Relationships in A: {len(rels_a)}")
    for r in rels_a:
        print(f" - {r['source_entity_id']} -> {r['target_entity_id']} ({r['type']})")
        
    print("\n--- RUNNING TEST B ---")
    case_b_id = f'test_case_B_{uuid.uuid4().hex[:8]}'
    await db["cases"].insert_one({"_id": case_b_id, "name": "Test Case B"})
    ev_b_result = await db['evidence'].insert_one({'case_id': case_b_id, 'title': 'testB.docx', 'source_type': 'DOCX'})
    ev_b_id = str(ev_b_result.inserted_id)
    
    with open("testB.docx", "rb") as f:
        content_B = f.read()
        
    await process_file_background(db, case_b_id, content_B, "testB.docx", current_user, ev_b_id)
    
    ents_b = await db["entities"].find({"case_id": case_b_id}).to_list(length=100)
    rels_b = await db["relationships"].find({"case_id": case_b_id}).to_list(length=100)
    ev_b = await db["evidence"].find_one({"_id": ev_b_result.inserted_id})
    
    print(f"Text B Length: {len(ev_b.get('raw_content', ''))}")
    print(f"Entities in B: {len(ents_b)}")
    for e in ents_b:
        print(f" - {e['name']} ({e['type']})")
    print(f"Relationships in B: {len(rels_b)}")
    for r in rels_b:
        print(f" - {r['source_entity_id']} -> {r['target_entity_id']} ({r['type']})")
        
    os.remove("testA.docx")
    os.remove("testB.docx")
    
if __name__ == '__main__':
    asyncio.run(main())
